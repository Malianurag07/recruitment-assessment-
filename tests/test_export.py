"""Excel and CSV export, including protection against spreadsheet formula injection from hostile resumes."""
import io
import json

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import deps
from app.main import app
from conftest import _judge, add_candidate


@pytest.fixture
def client(pool):
    conn, job_id = pool
    app.dependency_overrides[deps.get_db] = lambda: conn
    with TestClient(app) as c:
        c.job_id, c.conn = job_id, conn
        yield c
    app.dependency_overrides.clear()


def test_excel_export_has_ranking_and_a_coloured_skills_matrix(client):
    r = client.get(f"/api/jobs/{client.job_id}/export.xlsx")
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"] and "attachment" in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Ranking", "Skills matrix"]
    ranking = list(wb["Ranking"].values)
    assert ranking[0][:3] == ("Rank", "Name", "Email") and ranking[1][1] == "Jeevan Raj" and len(ranking) == 6
    matrix = wb["Skills matrix"]
    header = [c.value for c in matrix[1]]
    assert header[:2] == ["Candidate", "Score"] and "Python" in header and any("bonus" in str(h) for h in header)
    assert matrix.cell(2, header.index("Python") + 1).fill.fgColor.rgb.endswith("C6EFCE")       # green = solid skill
    assert client.get("/api/jobs/999/export.xlsx").status_code == 404


def test_excel_export_for_a_job_with_no_candidates_still_works(client):
    client.conn.execute("INSERT INTO job_descriptions (title) VALUES ('Empty')")
    client.conn.commit()
    r = client.get("/api/jobs/2/export.xlsx")
    assert r.status_code == 200 and len(list(load_workbook(io.BytesIO(r.content))["Ranking"].values)) == 1


def test_formula_injection_in_a_candidate_name_is_neutralised(client):
    add_candidate(client.conn, client.job_id, '=HYPERLINK("http://evil.example","click")', "evil@x.com", "9444444444", ["Python"], 0)
    csv_text = client.get(f"/api/jobs/{client.job_id}/export.csv").text
    assert "'=HYPERLINK" in csv_text and ",=HYPERLINK" not in csv_text
    wb = load_workbook(io.BytesIO(client.get(f"/api/jobs/{client.job_id}/export.xlsx").content))
    names = [row[1] for row in wb["Ranking"].values]
    assert any(str(n).startswith("'=HYPERLINK") for n in names) and not any(str(n).startswith("=") for n in names)
    assert not any(str(row[0]).startswith("=") for row in wb["Skills matrix"].values)
