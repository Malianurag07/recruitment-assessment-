"""Shared pieces for the QA suites (qa_offline.py, qa_live.py): a case recorder and resume generators.

Every case has an id, a title, what we expect, what actually happened, and PASS / FAIL. A FAIL is a finding, not a crash:
the suite keeps going and the report lists it. Results are written to docs/qa_results/<suite>.json.
"""
import io
import json
import os
import sys
import tempfile
import time
import traceback
from dataclasses import asdict, dataclass
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "docs" / "qa_results"
sys.path.insert(0, str(ROOT))


def use_throwaway_db(prefix: str) -> Path:
    """Must be called before importing anything from app: points DATABASE_PATH at a temporary file."""
    tmp = Path(tempfile.mkdtemp(prefix=prefix))
    os.environ["DATABASE_PATH"] = str(tmp / "qa.db")
    return tmp


@dataclass
class Case:
    scenario: str
    id: str
    title: str
    expected: str
    actual: str
    passed: bool
    seconds: float


class Recorder:
    def __init__(self, suite: str):
        self.suite, self.cases = suite, []

    def run(self, scenario: str, cid: str, title: str, expected: str, fn):
        """fn() returns (passed, actual_text). An exception is recorded as a FAIL with the error, never raised."""
        t = time.perf_counter()
        try:
            passed, actual = fn()
        except Exception as e:  # noqa: BLE001 - a test harness must survive everything
            passed, actual = False, f"EXCEPTION {type(e).__name__}: {e}\n{traceback.format_exc(limit=2)}"
        c = Case(scenario, cid, title, expected, str(actual), bool(passed), round(time.perf_counter() - t, 2))
        self.cases.append(c)
        print(f"  [{'PASS' if c.passed else 'FAIL'}] {cid} {title}  ({c.seconds}s)" + ("" if c.passed else f"\n         got: {c.actual[:300]}"), flush=True)
        return c

    def save(self, extra: dict | None = None):
        OUT.mkdir(parents=True, exist_ok=True)
        data = {"suite": self.suite, "extra": extra or {}, "cases": [asdict(c) for c in self.cases]}
        (OUT / f"{self.suite}.json").write_text(json.dumps(data, indent=1, ensure_ascii=False), encoding="utf-8")
        by = {}
        for c in self.cases:
            by.setdefault(c.scenario, [0, 0])[0 if c.passed else 1] += 1
        print("\n==== SUMMARY", self.suite)
        for s, (p, f) in by.items():
            print(f"  {s:45} {p:3} pass  {f:3} fail")
        print(f"  TOTAL {sum(c.passed for c in self.cases)}/{len(self.cases)} passed")


# ------------------------------------------------------------------ resume generators
def pdf_single(name, email, phone, body, white_text: str = "") -> bytes:
    import pymupdf
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((50, 60), name, fontsize=20)
    page.insert_text((50, 82), f"{email} | {phone}", fontsize=10)
    page.insert_textbox(pymupdf.Rect(50, 100, 545, 780), body, fontsize=10)
    if white_text:                                              # invisible to a human, visible to text extraction
        page.insert_textbox(pymupdf.Rect(50, 785, 545, 830), white_text, fontsize=4, color=(1, 1, 1))
    return doc.tobytes()


def pdf_two_column(name, email, phone, sidebar, main) -> bytes:
    import pymupdf
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((40, 50), name.upper(), fontsize=22)
    page.insert_textbox(pymupdf.Rect(40, 80, 200, 800), f"CONTACT\n{email}\n{phone}\n\n{sidebar}", fontsize=10)
    page.insert_textbox(pymupdf.Rect(230, 80, 560, 800), main, fontsize=10)
    return doc.tobytes()


def docx_bytes(paragraphs: list[str], table: list[list[str]] | None = None, header: str = "", footer: str = "") -> bytes:
    import docx
    d = docx.Document()
    if header:
        d.sections[0].header.paragraphs[0].text = header
    if footer:
        d.sections[0].footer.paragraphs[0].text = footer
    for p in paragraphs:
        d.add_paragraph(p)
    if table:
        t = d.add_table(rows=len(table), cols=len(table[0]))
        for i, row in enumerate(table):
            for j, v in enumerate(row):
                t.cell(i, j).text = v
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def timed(fn, *a, **k):
    t = time.perf_counter()
    r = fn(*a, **k)
    return r, time.perf_counter() - t
