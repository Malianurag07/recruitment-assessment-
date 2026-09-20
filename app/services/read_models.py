"""Read-side shapes for the web UI: candidate cards, job listings, CSV export, and a safe table viewer."""
import csv
import io
import json
import sqlite3

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.llm.query_tools import Ctx, load_candidates
from app.services import candidate_service as svc
from app.services.skill_normalizer import load_aliases

VIEWABLE_TABLES = ("candidates", "job_descriptions", "job_required_skills", "applications", "application_skills",
                   "analysis_results", "verification_log", "skill_aliases", "chat_history")


def list_jobs(conn: sqlite3.Connection) -> list[dict]:
    jobs = []
    for r in conn.execute("SELECT id, title, created_at FROM job_descriptions ORDER BY id DESC"):
        counts = conn.execute(
            """SELECT SUM(application_status='active') AS active, SUM(application_status='pending_choice') AS pending
               FROM applications WHERE job_description_id=?""", (r["id"],)).fetchone()
        scored = conn.execute("SELECT COUNT(*) FROM analysis_results WHERE job_description_id=?", (r["id"],)).fetchone()[0]
        jobs.append({"id": r["id"], "title": r["title"], "created_at": r["created_at"], "scored": scored,
                     "pending": counts["pending"] or 0})
    return jobs


def job_detail(conn: sqlite3.Connection, job_id: int) -> dict | None:
    job = svc.get_job(conn, job_id)
    if job is None:
        return None
    ctx = Ctx(conn, job_id, load_aliases(conn), job)
    cands = load_candidates(ctx)
    by = {"Shortlist": 0, "Consider": 0, "Reject": 0}
    for c in cands:
        by[c["recommendation"]] = by.get(c["recommendation"], 0) + 1
    review = conn.execute(
        "SELECT COUNT(*) FROM applications WHERE job_description_id=? AND application_status='active' "
        "AND (extraction_status!='ok' OR verification_status='needs_review')", (job_id,)).fetchone()[0]
    failed = conn.execute("SELECT COUNT(*) FROM analysis_results WHERE job_description_id=? AND llm_status='unavailable'",
                          (job_id,)).fetchone()[0]
    return {"id": job_id, **job.model_dump(), "scored": len(cands), "by_recommendation": by,
            "average_score": round(sum(c["score"] for c in cands) / len(cands), 1) if cands else None,
            "pending_choices": len(svc.pending_conflicts(conn, job_id)), "needs_review": review, "failed_analyses": failed}


def candidate_cards(conn: sqlite3.Connection, job_id: int) -> list[dict]:
    job = svc.get_job(conn, job_id)
    if job is None:
        return []
    cards = []
    for c in load_candidates(Ctx(conn, job_id, load_aliases(conn), job)):
        log = conn.execute("SELECT field, old_value, new_value, evidence_quote FROM verification_log WHERE application_id=? ORDER BY id",
                           (c["application_id"],)).fetchall()
        status = conn.execute("SELECT llm_status FROM analysis_results WHERE application_id=?", (c["application_id"],)).fetchone()
        req = [b for b in c["breakdown"] if b["importance"] == "required"]
        all_skills = [dict(r) for r in conn.execute(
            "SELECT skill_name AS skill, raw_skill AS raw, level FROM application_skills WHERE application_id=? ORDER BY id",
            (c["application_id"],))]
        cards.append({
            "application_id": c["application_id"], "rank": c["rank"], "name": c["name"], "email": c["email"], "phone": c["phone"],
            "file": c["file"], "score": c["score"], "recommendation": c["recommendation"], "components": c["components"],
            "required_matched": sum(b["status"] != "missing" for b in req), "required_total": len(req),
            "skills": c["breakdown"], "summary": c["summary"], "strengths": c["strengths"], "weaknesses": c["weaknesses"],
            "interview_questions": c["questions"], "experience_years": c["exp_years"], "internships": c["internships"],
            "soft_skills": c["soft_skills"], "projects": c["projects"], "certifications": c["certifications"],
            "education": c["education"], "verification": c["verification"],
            "all_skills": all_skills, "verification_log": [dict(r) for r in log], "llm_status": status["llm_status"] if status else None})
    return cards


def needs_review(conn: sqlite3.Connection, job_id: int) -> list[dict]:
    rows = conn.execute(
        """SELECT a.id AS application_id, c.name, c.email, a.resume_filename, a.extraction_status, a.verification_status
           FROM applications a JOIN candidates c ON c.id = a.candidate_id
           WHERE a.job_description_id=? AND a.application_status='active'
             AND (a.extraction_status!='ok' OR a.verification_status='needs_review')""", (job_id,)).fetchall()
    return [dict(r) for r in rows]


_FORMULA_STARTS = ("=", "+", "-", "@", "\t", "\r")


def safe_cell(value):
    """Neutralise spreadsheet formula injection. Names and skills come from untrusted resumes: a name such as
    '=HYPERLINK(...)' must open as text, not run as a formula. A leading apostrophe forces text in Excel and Sheets."""
    if isinstance(value, str) and value.startswith(_FORMULA_STARTS):
        return "'" + value
    return value


def ranking_csv(conn: sqlite3.Connection, job_id: int) -> str:
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(["rank", "name", "email", "phone", "score", "recommendation", "required_skills_matched", "experience_years",
                "matching_skills", "missing_skills", "resume_file"])
    for c in candidate_cards(conn, job_id):
        have = [s["skill"] for s in c["skills"] if s["status"] != "missing"]
        miss = [s["skill"] for s in c["skills"] if s["status"] == "missing"]
        row = [c["rank"], c["name"], c["email"], c["phone"], c["score"], c["recommendation"],
               f"{c['required_matched']}/{c['required_total']}", c["experience_years"], "; ".join(have), "; ".join(miss), c["file"]]
        w.writerow([safe_cell(v) for v in row])
    return out.getvalue()


# Excel fills for the recommendation and for the skill badge colours (same meaning as in the web UI)
_REC_FILL = {"Shortlist": "C6EFCE", "Consider": "FFEB9C", "Reject": "FFC7CE"}
_BADGE_FILL = {"green": "C6EFCE", "yellow": "FFEB9C", "orange": "F8CBAD", "red": "FFC7CE"}


def ranking_xlsx(conn: sqlite3.Connection, job_id: int) -> bytes:
    """Two-sheet workbook: the ranking, and a candidate x skill matrix coloured like the UI's skill badges."""
    job = svc.get_job(conn, job_id)
    cards = candidate_cards(conn, job_id)
    wb = Workbook()
    ws = wb.active
    ws.title = "Ranking"
    ws.append(["Rank", "Name", "Email", "Phone", "Score", "Recommendation", "Required skills", "Experience (yrs)",
               "Internships", "Summary", "Strengths", "Gaps", "Resume file"])
    for c in cards:
        ws.append([safe_cell(v) for v in [
            c["rank"], c["name"], c["email"], c["phone"], c["score"], c["recommendation"],
            f"{c['required_matched']}/{c['required_total']}", c["experience_years"], len(c["internships"]), c["summary"],
            "\n".join(c["strengths"]), "\n".join(c["weaknesses"]), c["file"]]])
        ws.cell(ws.max_row, 6).fill = PatternFill("solid", fgColor=_REC_FILL.get(c["recommendation"], "FFFFFF"))
    for i, width in enumerate([6, 26, 30, 16, 8, 15, 15, 16, 12, 60, 50, 50, 28], 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    sk = wb.create_sheet("Skills matrix")
    skills = ([(s, "required") for s in job.required_skills] + [(s, "preferred") for s in job.preferred_skills]) if job else []
    sk.append(["Candidate", "Score"] + [f"{name}{' (bonus)' if imp == 'preferred' else ''}" for name, imp in skills])
    for c in cards:
        by_name = {s["skill"]: s for s in c["skills"]}
        sk.append([safe_cell(c["name"]), c["score"]] + [by_name[n]["status"] if n in by_name else "" for n, _ in skills])
        for j, (n, _) in enumerate(skills, 3):
            if n in by_name:
                sk.cell(sk.max_row, j).fill = PatternFill("solid", fgColor=_BADGE_FILL[by_name[n]["colour"]])
    sk.column_dimensions["A"].width = 26
    for j in range(3, 3 + len(skills)):
        sk.column_dimensions[get_column_letter(j)].width = 16

    for sheet in (ws, sk):
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def table_rows(conn: sqlite3.Connection, table: str, limit: int = 100) -> dict:
    """Read-only table viewer for the demo. The table name is checked against a fixed whitelist, never interpolated blindly."""
    if table not in VIEWABLE_TABLES:
        raise KeyError(table)
    total = conn.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]          # noqa: S608 (whitelisted above)
    cur = conn.execute(f"SELECT * FROM {table} ORDER BY 1 DESC LIMIT ?", (max(1, min(limit, 500)),))  # noqa: S608
    cols = [d[0] for d in cur.description]
    rows = [[(str(v)[:160] + "…") if isinstance(v, str) and len(v) > 160 else v for v in r] for r in cur.fetchall()]
    return {"table": table, "columns": cols, "rows": rows, "total": total}


def table_counts(conn: sqlite3.Connection) -> dict:
    return {t: conn.execute(f"SELECT COUNT(*) FROM {t}").fetchone()[0] for t in VIEWABLE_TABLES}  # noqa: S608
